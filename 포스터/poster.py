"""Mars-Fit 포스터 (900 x 1200 mm)

연구실 틀(`이거 수정해서 만들거임.pptx`)의 색·글꼴·2단 구조를 그대로 따르되,
연구 포스터의 「서론-방법-결과」 대신 서비스 포스터로 짠다.

섹션 높이는 내용을 다 그린 뒤에 정해진다(section_open/close). 패널을 먼저
그리고 높이를 어림하면 빈 칸이 크게 남는다.

같은 좌표를 HTML 로도 뽑아 브라우저로 찍어본다(preview.html) — pptx 렌더러가
없어서 글이 상자를 넘치는지 눈으로 볼 길이 그것뿐이다.
"""
import io
import os
from pptx import Presentation
from pptx.util import Mm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

S = os.path.dirname(os.path.abspath(__file__))
# 화면 사진. 저장소에서는 `포스터/화면/`, 작업 중에는 스크래치패드의 `crop/`.
CROP = f"{S}/화면" if os.path.isdir(f"{S}/화면") else f"{S}/crop"

# ── 디자인 토큰 ────────────────────────────────────────────────
NAVY   = RGBColor(0x00, 0x36, 0x70)
ORANGE = RGBColor(0xCB, 0x6B, 0x3D)
INK    = RGBColor(0x2B, 0x2B, 0x2B)
MUTED  = RGBColor(0x6B, 0x64, 0x59)
PAPER  = RGBColor(0xF5, 0xF3, 0xED)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LINE   = RGBColor(0xD8, 0xD2, 0xC4)
TINT   = RGBColor(0xEC, 0xF1, 0xF7)
GREEN  = RGBColor(0x0E, 0x7A, 0x5A)
GOLD   = RGBColor(0xF2, 0xB8, 0x6A)
PALE   = RGBColor(0xC7, 0xD4, 0xE6)

H_FONT, B_FONT = '나눔스퀘어_ac Bold', '맑은 고딕'

# 글자 배율. 연구실 틀의 본문이 25pt 이고, 예시 포스터 10개를 900x1200 으로
# 환산하면 22~28pt(중앙값 24)이다. 처음에 19~21pt 로 잡아 20% 작았다.
SCALE = 1.22

W, H = 900, 1200
HEAD_H, FOOT_H = 150, 26
L_X, R_X, COL_W = 12, 460, 428
TOP, BOT = HEAD_H + 14, H - FOOT_H - 14
PAD = 14                                  # 패널 안쪽 여백

prs = Presentation()
prs.slide_width, prs.slide_height = Mm(W), Mm(H)
slide = prs.slides.add_slide(prs.slide_layouts[6])
SH = slide.shapes
OPS = []


# ── 기본 도형 ──────────────────────────────────────────────────
def box(x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, adj=None):
    s = SH.add_shape(shape, Mm(x), Mm(y), Mm(w), Mm(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    if adj is not None:
        try: s.adjustments[0] = adj
        except Exception: pass
    s.text_frame.clear()
    kind = ('oval' if shape == MSO_SHAPE.OVAL
            else 'round' if shape == MSO_SHAPE.ROUNDED_RECTANGLE else 'rect')
    OPS.append(dict(k='box', x=x, y=y, w=w, h=h,
                    fill=fill and str(fill), line=line and str(line), lw=lw,
                    shape=kind, adj=adj))
    return s, OPS[-1]


def fill_shape(shp, rec, lines, size, color=WHITE, font=B_FONT, bold=True,
               align=PP_ALIGN.CENTER):
    """도형 안에 글을 넣고 미리보기에도 같이 기록한다."""
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Mm(2)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        r.font.name, r.font.size, r.font.bold = font, Pt(size * SCALE), bold
        r.font.color.rgb = color
    rec.update(txt=lines, tsize=size * SCALE, tcolor=str(color), tbold=bold)


def text(x, y, w, h, lines, size=22, color=INK, font=B_FONT, bold=False,
         align=PP_ALIGN.LEFT, space=0.30, line_sp=1.24):
    tb = SH.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    size = size * SCALE
    if isinstance(lines, str):
        lines = [lines]
    rec = []
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_sp
        if i:
            p.space_before = Pt(size * space)
        parts = item if isinstance(item, list) else [item]
        rrec = []
        for part in parts:
            t, o = (part, {}) if isinstance(part, str) else part
            r = p.add_run(); r.text = t
            r.font.name = o.get('font', font)
            r.font.size = Pt(o.get('size', size / SCALE) * SCALE)
            r.font.bold = o.get('bold', bold)
            r.font.color.rgb = o.get('color', color)
            oo = {k: (str(v) if k == 'color' else v) for k, v in o.items()}
            if 'size' in oo: oo['size'] = oo['size'] * SCALE
            rrec.append((t, oo))
        rec.append(rrec)
    OPS.append(dict(k='text', x=x, y=y, w=w, h=h, rec=rec, size=size,
                    color=str(color), font=font, bold=bold,
                    align=str(align).split('.')[-1].split(' ')[0],
                    space=space, line_sp=line_sp))
    return tb


def pic(path, x, y, w=None, h=None):
    p = SH.add_picture(path, Mm(x), Mm(y), Mm(w) if w else None, Mm(h) if h else None)
    pw, ph = Emu(p.width).inches * 25.4, Emu(p.height).inches * 25.4
    OPS.append(dict(k='pic', x=x, y=y, w=pw, h=ph, path=path))
    return pw, ph


def shot(path, x, y, h):
    """휴대폰 화면. 높이를 주면 비율대로 폭이 정해지고 테두리를 두른다."""
    w, _ = pic(path, x, y, h=h)
    box(x, y, w, h, fill=None, line=LINE, lw=0.9)
    return w


# ── 높이를 나중에 정하는 섹션 ──────────────────────────────────
TAB_H = 21
SECTIONS = []


def section_open(x, y, w, title):
    """내용 시작 y 와 손잡이를 돌려준다. 패널은 close 에서 뒤에 깔린다."""
    return y + TAB_H + 9, (x, y, w, title, len(SH._spTree), len(OPS))


def section_close(handle, cy):
    x, y, w, title, sp_idx, ops_idx = handle
    panel_y = y + TAB_H * 0.55
    panel, prec = box(x, panel_y, w, (cy + PAD) - panel_y, fill=WHITE, line=LINE,
                      lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.030)
    tw = min(w * 0.54, 216)
    tab, trec = box(x + 8, y, tw, TAB_H, fill=NAVY,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.42)
    fill_shape(tab, trec, [title], 30, font=H_FONT)
    for shp in (tab, panel):                       # 그린 내용 뒤로 보낸다
        el = shp._element
        el.getparent().remove(el)
        SH._spTree.insert(sp_idx, el)
    for rec in (trec, prec):
        OPS.remove(rec); OPS.insert(ops_idx, rec)
    end = cy + PAD
    SECTIONS.append((title, x, y, end - y))
    return end


def bullet(x, y, w, items, size=21, gap=1.6):
    cy = y
    for it in items:
        parts = it if isinstance(it, list) else [it]
        box(x + 1.4, cy + 3.4, 3.0, 3.0, fill=ORANGE, shape=MSO_SHAPE.OVAL)
        n = sum(len(t if isinstance(t, str) else t[0]) for t in parts)
        s2 = size * SCALE
        per = int((w - 9) / (s2 * 0.372))
        hh = max(1, -(-n // per)) * s2 * 0.437
        text(x + 9, cy, w - 9, hh, [parts], size=size)
        cy += hh + gap
    return cy - gap


def formula(x, y, w, lines, size=23):
    hh = len(lines) * size * SCALE * 0.52 + 10
    box(x, y, w, hh, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.10)
    text(x + 7, y + 5, w - 14, hh - 10, lines, size=size, font='Consolas',
         color=NAVY, bold=True, align=PP_ALIGN.CENTER, line_sp=1.32)
    return y + hh


def table(x, y, w, widths, head, rows, size=19, head_size=16):
    """판정 표. 셀은 str 또는 (str, {서식}).

    글줄 셋을 나란히 늘어놓으면 「판정 이름 / 뜻 / 값 / 건수」가 섞여 읽힌다.
    칸을 그어야 무엇이 무엇에 대응하는지가 한눈에 잡힌다.
    """
    HEAD_H, ROW_H = 15, 21
    total = sum(widths)
    xs, acc = [], x
    for cw in widths:
        xs.append(acc); acc += cw / total * w
    xs.append(x + w)

    box(x, y, w, HEAD_H, fill=NAVY, line=None,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.14)
    for i, h in enumerate(head):
        text(xs[i] + 5, y + 3.6, xs[i + 1] - xs[i] - 10, 10, h, size=head_size,
             color=RGBColor(0xC7, 0xD4, 0xE6),
             align=PP_ALIGN.LEFT if i == 1 else PP_ALIGN.CENTER)

    cy = y + HEAD_H
    for r, row in enumerate(rows):
        if r:
            box(x + 4, cy, w - 8, 0.4, fill=LINE)
        for i, cell in enumerate(row):
            t, o = (cell, {}) if isinstance(cell, str) else cell
            cw = xs[i + 1] - xs[i]
            if o.get('pill'):
                b, rec = box(xs[i] + 6, cy + 4.0, cw - 12, 13, fill=o['pill'],
                             shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.40)
                fill_shape(b, rec, [t], 16)
            else:
                text(xs[i] + 5, cy + 5.0, cw - 10, 12, t, size=o.get('size', size),
                     color=o.get('color', INK), bold=o.get('bold', False),
                     align=PP_ALIGN.LEFT if i == 1 else PP_ALIGN.CENTER)
        cy += ROW_H
    box(x, y, w, cy - y, fill=None, line=LINE, lw=1.0,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.045)
    return cy


def para(x, y, w, lines, size=20, **kw):
    """문단 묶음. 끝 y 를 돌려준다.

    **줄 수는 문단마다 따로 센다.** 전체 글자수로 한 번에 세면 짧은 줄 셋을
    한 줄로 보고 높이를 1/3 로 잡아서, 다음 블록이 그 위에 겹쳐 그려진다.
    """
    s2 = size * SCALE
    per = int(w / (s2 * 0.372))
    rows = 0
    for item in lines:
        n = sum(len(t if isinstance(t, str) else t[0])
                for t in (item if isinstance(item, list) else [item]))
        rows += max(1, -(-n // per))
    hh = rows * s2 * 0.437 + (len(lines) - 1) * s2 * kw.get('space', 0.30) * 0.353
    text(x, y, w, hh, lines, size=size, **kw)
    return y + hh


# ══════════════════════════════════════════════════════════════
# 배경 · 머리 · 꼬리
# ══════════════════════════════════════════════════════════════
box(0, 0, W, H, fill=PAPER)
box(0, 0, W, HEAD_H, fill=NAVY)
box(0, H - FOOT_H, W, FOOT_H, fill=NAVY)

text(48, 21, 560, 22, '2026 AI화성 챌린지 in 수원대학교   ·   최우수상',
     size=25, color=GOLD, bold=True, font=H_FONT)
text(48, 47, 700, 46, 'Mars-Fit — 화성시 소상공인 AI 경영동행 서비스',
     size=50, color=WHITE, bold=True, font=H_FONT)
text(48, 101, 700, 24,
     '흩어진 지원사업을 내 조건으로 판정하고, 서류 준비와 신청까지 데려다주는 웹 서비스',
     size=24, color=PALE)
text(48, 127, 700, 18,
     [[('임대윤 · 전서희 · 전성현', {'bold': True}),
       ('     수원대학교 데이터과학부', {'color': RGBColor(0xA9, 0xBC, 0xD6)})]],
     size=21, color=WHITE)

box(W - 146, 20, 110, 110, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06)
pic(f"{CROP}/qr.png", W - 140, 26, w=98)
text(W - 146, 132, 110, 14, '휴대폰으로 열어보세요', size=15, color=PALE,
     align=PP_ALIGN.CENTER)

text(0, H - FOOT_H + 7, W, 16,
     'https://ai-hwaseong-ten.vercel.app     ·     React 18 + Vite  ·  Vercel Serverless  ·  Supabase  ·  제공 서버 cron',
     size=16, color=RGBColor(0xB9, 0xCB, 0xE2), align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════
# 좌단
# ══════════════════════════════════════════════════════════════
y = TOP
IN_X, IN_W = L_X + PAD, COL_W - PAD * 2

# ① 왜 만들었나
cy, hd = section_open(L_X, y, COL_W, '왜 만들었나')
cy = bullet(IN_X, cy, IN_W, [
    [('지원사업이 한곳에 모여 있지 않다.', {'bold': True, 'color': NAVY}),
     ' 중앙(기업마당)과 화성시청 고시공고를 따로 뒤져야 한다.'],
    [('화성시가 직접 하는 사업은 중앙 목록에 안 올라온다.', {'bold': True, 'color': NAVY}),
     ' 소상공인 자금지원(특례보증 5천만원·이차보전 2%)과 저신용 미소금융 이자지원이 '
     '그렇다 — 지금 열려 있는데 중앙에서는 보이지 않는다.'],
    [('자격 요건이 공문서 언어다.', {'bold': True, 'color': NAVY}),
     ' 「내가 되는지」를 사장님이 스스로 판단하기 어렵다.'],
]) + 7
cy = para(IN_X, cy, IN_W, ['그리고 마감은 되돌릴 수 없다.'],
          size=22, color=ORANGE, bold=True)
y = section_close(hd, cy) + 12

# ② 사장님이 겪는 흐름
cy, hd = section_open(L_X, y, COL_W, '사장님이 겪는 흐름')
STEPS = ['등록증 사진 1장', '조건 자동 매칭', '조건별 판정', '서류 발급 안내', '접수 · 알림']
# 칸은 글자에 맞춰 좁히고 화살표를 키운다. 헐렁한 칸에 작은 화살표면
# 「다음 단계로 간다」가 안 읽힌다.
sgap = 8.0
sw = (IN_W - sgap * (len(STEPS) - 1)) / len(STEPS)
sx = IN_X
for i, st in enumerate(STEPS):
    b, rec = box(sx, cy, sw, 22, fill=NAVY if i == 0 else WHITE,
                 line=None if i == 0 else NAVY, lw=1.2,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.34)
    fill_shape(b, rec, [st], 16, color=WHITE if i == 0 else NAVY)
    if i < len(STEPS) - 1:
        text(sx + sw, cy - 0.5, sgap, 23, '›', size=30, color=ORANGE, bold=True,
             align=PP_ALIGN.CENTER)
    sx += sw + sgap
cy += 32

w1 = shot(f"{CROP}/onboard.png", IN_X + 22, cy, 232)
w2 = shot(f"{CROP}/home_list.png", IN_X + 22 + w1 + 22, cy, 232)
text(IN_X + 22, cy + 235, w1, 14, '등록증을 올리면 자동으로 읽는다',
     size=15, color=MUTED, align=PP_ALIGN.CENTER)
text(IN_X + 22 + w1 + 22, cy + 235, w2, 14, '내 조건으로 걸러진 목록 · 긴급 마감',
     size=15, color=MUTED, align=PP_ALIGN.CENTER)
cy += 256
cy = para(IN_X, cy, IN_W, [
    [('사진 한 장이면 시작된다.', {'bold': True, 'color': NAVY}),
     ' 사업자등록증을 찍으면 상호·업종·개업일을 읽어 프로필을 채운다. '
     '사진은 메모리에서만 읽고 저장하지 않으며 등록증 전문도 돌려주지 않는다.'],
], size=20)
y = section_close(hd, cy) + 12

# ③ 조건 판정 매칭 엔진
cy, hd = section_open(L_X, y, COL_W, '조건 판정 매칭 엔진')
cy = para(IN_X, cy, IN_W, [
    [('공고 한 건을 12개 조건축으로 쪼개', {'bold': True, 'color': NAVY}),
     ' 사장님 프로필과 하나씩 맞춰본다. 「지원사업 목록」이 아니라 「판정」을 준다.'],
], size=20) + 8
box(IN_X, cy, IN_W, 32, fill=TINT, line=None,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.14)
text(IN_X + 6, cy + 5, IN_W - 12, 24,
     ['지역 · 업종 · 나이 · 사업 상태 · 사업자 형태 · 창업 경험 · 소득·자산 구간',
      '운영 기간 · 연 매출 · 혼인 상태 · 부모 동거 여부 · 지원 자격'],
     size=17, color=NAVY, align=PP_ALIGN.CENTER, line_sp=1.30, space=0.10)
cy += 40
cy = formula(IN_X, cy, IN_W, ['매칭 점수 =  Σ( wᵢ × vᵢ )  ÷  Σ wᵢ  × 100'], size=24) + 10
cy = table(IN_X, cy, IN_W, [78, 196, 52, 74],
           ['판정', '무슨 뜻인가', '조건값 vᵢ', '오늘 74건'],
           [[('신청가능', {'pill': GREEN}), '조건이 모두 맞는다',
             ('1.0', {'bold': True, 'color': GREEN}), '31건'],
            [('확인필요', {'pill': ORANGE}), '공고문만으로는 판단이 서지 않는다',
             ('0.5', {'bold': True, 'color': ORANGE}), '27건'],
            [('대상아님', {'pill': MUTED}), '맞지 않는 조건이 분명히 있다',
             ('0.0', {'bold': True, 'color': MUTED}), '16건']]) + 9


sw2 = shot(f"{CROP}/judge.png", IN_X, cy, 150)
tx, tw = IN_X + sw2 + 16, IN_W - sw2 - 16
ty = para(tx, cy - 1, tw,
          [[('「확인필요」를 지우지 않는다.', {'bold': True, 'color': NAVY})]], size=20)
ty = para(tx, ty + 4, tw, [
    '모르는 것을 아는 척도, 모르는 척도 하지 않는다. 업종이 못 박힌 공고는 '
    '대상아님이 아니라 확인필요로 내린다 — 제조업을 겸하는 사장님이 있기 때문이다.',
], size=18)
ty = para(tx, ty + 11, tw,
          [[('조건마다 무게가 다르다', {'bold': True, 'color': NAVY, 'size': 20})]], size=20)
ty = para(tx, ty + 4, tw, [
    '「소상공인이면 누구나」에는 0.7점만 준다. 자격은 되지만 나에게 맞는 사업은 '
    '아니다. 1.0 을 주면 요건이 그것뿐인 공고가 전부 100점으로 묶여 순위가 사라진다.',
], size=18)
ty = para(tx, ty + 11, tw,
          [[('「불충족」은 확실할 때만 쓴다', {'bold': True, 'color': NAVY, 'size': 20})]],
          size=20)
ty = para(tx, ty + 4, tw, [
    '하나라도 있으면 그 공고는 대상아님이 되어 목록에서 사라진다. '
    '애매한 것은 여기 넣지 않고 확인필요로 남긴다.'], size=18)
cy = max(cy + 150, ty) + 10

cy = para(IN_X, cy, IN_W, [
    [('오늘 74건에서 조건 309개를 판정했다', {'bold': True, 'color': NAVY}),
     ' — 충족 247 · 확인필요 46 · 불충족 16. 공고마다 서류도 같이 붙는다'
     '(195개 · 52종).'],
    [('줄 세우기는 접수중 먼저, 그 안에서 점수순이다.', {'bold': True, 'color': NAVY}),
     ' 판정을 점수보다 먼저 보면, 화성시 사업이 서류 한 줄로 확인필요가 되는 순간 '
     '전국 공고 수십 건 아래로 밀린다. 확인필요는 이미 0.5 로 반영돼 있다.'],
], size=19, space=0.42)
y = section_close(hd, cy)


# ══════════════════════════════════════════════════════════════
# 우단
# ══════════════════════════════════════════════════════════════
y = TOP
IN_X, IN_W = R_X + PAD, COL_W - PAD * 2

# ④ 서류는 계산해서 알려준다  (조건 판정 바로 다음에 읽히게 우단 맨 위)
cy, hd = section_open(R_X, y, COL_W, '서류는 계산해서 알려준다')
cy = para(IN_X, cy, IN_W, [
    [('찾아주는 데서 끝내지 않는다.', {'bold': True, 'color': NAVY}),
     ' 판정이 끝나면 「그래서 뭘 내야 하나」가 남는다. 공고마다 낼 서류가 다르고, '
     '오늘 74건에 서류가 195개 · 52종 붙어 있다.'],
], size=20) + 10

sw = shot(f"{CROP}/apply.png", IN_X, cy, 178)
tx, tw = IN_X + sw + 16, IN_W - sw - 16

ty = para(tx, cy - 1, tw,
          [[('세 단계로 계산한다', {'bold': True, 'color': NAVY, 'size': 21})]], size=21)
STEPS_D = [
    ('1', '공고문에서 뽑는다',
     '「제출서류·구비서류」 머리말 아래의 표와 목록을 읽는다. 쉼표로 묶인 줄은 '
     '나누고, 서류명 뒤에 붙은 발급처 안내는 떼어낸다.'),
    ('2', '프로필로 거른다',
     '서류마다 조건이 달려 있다 — 「직원이 있으면 4대보험 가입자명부」처럼. '
     '값을 모르면 일단 남긴다. 빠뜨리는 쪽이 더 나쁘다.'),
    ('3', '발급 정보를 붙인다',
     '발급 절차·수수료·소요시간을 얹고, 공통필수 · 조건부필수 · 해당시제출로 '
     '갈라 적는다.'),
]
for num, head, body in STEPS_D:
    b, rec = box(tx, ty + 4, 9, 9, fill=NAVY, shape=MSO_SHAPE.OVAL)
    fill_shape(b, rec, [num], 13)
    ty = para(tx + 13, ty + 1, tw - 13,
              [[(head, {'bold': True, 'color': NAVY, 'size': 19})]], size=19)
    ty = para(tx + 13, ty + 2, tw - 13, [body], size=17)
    ty += 6

ty = para(tx, ty + 2, tw,
          [[('그리고 그 기관으로 바로 보낸다', {'bold': True, 'color': NAVY, 'size': 21})]],
          size=21)
ty = para(tx, ty + 3, tw, [
    '「발급 절차 보기」를 누르면 사장님이 홈택스·정부24 를 따로 찾아 들어갈 일이 없다.',
], size=17)
cy = max(cy + 178, ty) + 10

# 실제로 연결되는 곳
box(IN_X, cy, IN_W, 42, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.10)
text(IN_X + 8, cy + 5, IN_W - 16, 14,
     '서류 26종에 붙여둔 발급처 — 온라인 절차 24종 · 바로 가는 링크 16종',
     size=15, color=MUTED, align=PP_ALIGN.CENTER)
SITES = [('홈택스', '7종'), ('정부24', '5종'), ('건강보험공단', '2종'),
         ('대법원 전자가족관계', '1종'), ('중소벤처24', '1종')]
sx, sgap = IN_X + 8, 5
bw = (IN_W - 16 - sgap * (len(SITES) - 1)) / len(SITES)
for nm, cnt in SITES:
    b, rec = box(sx, cy + 21, bw, 15, fill=WHITE, line=NAVY, lw=0.9,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.34)
    fill_shape(b, rec, [f'{nm} {cnt}'], 13, color=NAVY)
    sx += bw + sgap
cy += 46
cy = para(IN_X, cy, IN_W, [
    [('체크한 것은 공고마다 따로 남는다.', {'bold': True, 'color': NAVY}),
     ' 발급 3개월 이내 서류만 인정된다는 것도 같이 적는다.'],
], size=19)
y = section_close(hd, cy) + 12

# ④ 상권 추천 알고리즘
cy, hd = section_open(R_X, y, COL_W, '상권 추천 알고리즘')
cy = para(IN_X, cy, IN_W, [
    [('업종을 고르면 자리를 추천한다.', {'bold': True, 'color': NAVY}),
     ' 「고른 자리를 분석」하는 기존 서비스와 방향이 반대다. 화성시를 '
     '500m 격자로 나눠 칸마다 점수를 매긴다.'],
], size=20) + 8
cy = formula(IN_X, cy, IN_W,
             ['score = w_ft·유동인구 + w_dem·상권활성화 − w_comp·경쟁강도'], size=20) + 5
cy = formula(IN_X, cy, IN_W,
             ['유동인구 = 학교×220 + 카페×90 + 음식점×55',
              '+ 역 승하차×0.15 + 아파트 세대×0.08'], size=19) + 10

sw3 = shot(f"{CROP}/map.png", IN_X, cy, 168)
tx, tw = IN_X + sw3 + 16, IN_W - sw3 - 16
ty = para(tx, cy - 1, tw,
          [[('가중치를 사장님이 직접 조절한다', {'bold': True, 'color': NAVY, 'size': 20})]],
          size=20)
ty = para(tx, ty + 5, tw, [
    '유동인구 · 상권 활성화 · 경쟁 회피 세 축을 슬라이더로 옮기면 추천 자리가 '
    '그 자리에서 다시 계산된다. 기본값은 4 : 2 : 4.'], size=18)
ty = para(tx, ty + 11, tw,
          [[('숫자를 지어내지 않는다', {'bold': True, 'color': NAVY, 'size': 20})]], size=20)
ty = para(tx, ty + 5, tw, [
    '상가·학교·역·아파트는 모두 실측값이고, 유동인구만 그것을 가중합산한 '
    '추정치라 화면에 「예측」이라 적는다. 카드매출도 그 행정동 자료가 없으면 '
    '화성시 평균으로 떨어뜨리고 「이 자리 매출이 아닙니다」라고 밝힌다.'], size=18)
ty = para(tx, ty + 9, tw, [
    [('경쟁강도 기준도 업종마다 다르다 — ', {'bold': True, 'color': NAVY}),
     '같은 「10개」라도 카페는 붐비고 음식점은 한산하다. 카페 8·20, '
     '음식점 15·40, 소매업 10·30 이 경계다.']], size=18)
ty = para(tx, ty + 9, tw, [
    [('격자에 얹는 것 — ', {'bold': True, 'color': NAVY}),
     '상가 30,297 · 학교 188 · 지하철역 5 · 아파트 47개 행정동. '
     '반경은 300 · 500 · 1000m 로 바꿔 볼 수 있다.']], size=18)

cy = max(cy + 168, ty)
y = section_close(hd, cy) + 12

# ⑤ 매일 도는 공고 수집
cy, hd = section_open(R_X, y, COL_W, '매일 도는 공고 수집')
PIPE = [(['06:11', '수집'], NAVY), (['guard', '회귀 검증'], NAVY),
        (['커밋 · 푸시'], NAVY), (['GitHub Actions', '빌드 검사'], NAVY),
        (['Vercel', '자동 배포'], ORANGE)]
pgap = 9
pw = (IN_W - pgap * (len(PIPE) - 1)) / len(PIPE)
px = IN_X
for i, (t, c) in enumerate(PIPE):
    b, rec = box(px, cy, pw, 29, fill=c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.24)
    fill_shape(b, rec, t, 17)
    if i < len(PIPE) - 1:
        text(px + pw, cy + 2.5, pgap, 23, '›', size=30, color=ORANGE, bold=True,
             align=PP_ALIGN.CENTER)
    px += pw + pgap
cy += 35
cy = para(IN_X, cy, IN_W,
          ['기업마당 API + 화성시청 고시공고  ·  매일 1회'],
          size=16, color=MUTED) + 8
cy = bullet(IN_X, cy, IN_W, [
    [('수집만 제공 서버의 cron 이 맡는다.', {'bold': True, 'color': NAVY}),
     ' 기업마당 API 가 GitHub 러너 IP 에 응답하지 않아 스케줄을 껐다. '
     'Actions 는 대신 빌드 검사를 맡는다 — main 에 올리면 사람 손 없이 '
     'Vercel 이 그대로 실서비스에 올리기 때문이다.'],
    [('총량으로 검증하지 않는다.', {'bold': True, 'color': NAVY}),
     ' 마감된 공고가 빠지면 딸린 서류도 같이 빠져 「데이터가 깎였다」로 잘못 읽힌다. '
     '양쪽에 다 있는 공고만 골라 [서류·조건·본문]을 견준다.'],
    [('걸리면 커밋하지 않고 관리자 카카오톡으로 알린다.', {'bold': True, 'color': NAVY}),
     ' 로그 파일에만 적어두면 아무도 안 본다.'],
], size=19)
y = section_close(hd, cy) + 12

# ⑦ 지금 돌아가고 있다
cy, hd = section_open(R_X, y, COL_W, '지금 돌아가고 있다')
STATS = [('74', '공고 · 실시간'), ('12', '조건축'), ('18', '세무 항목'),
         ('43', '행정용어'), ('52', '서류 종류')]
bw = (IN_W - 4 * 8) / 5
bx = IN_X
for v, lab in STATS:
    box(bx, cy, bw, 40, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.14)
    text(bx, cy + 5, bw, 22, v, size=34, color=NAVY, bold=True, font=H_FONT,
         align=PP_ALIGN.CENTER)
    text(bx, cy + 24, bw, 16, lab, size=14, color=MUTED, align=PP_ALIGN.CENTER)
    bx += bw + 8
cy += 42
cy = para(IN_X, cy, IN_W,
          ['카카오 로그인 없이도 전부 쓸 수 있다  ·  기기 저장이 원본, 서버는 통로'],
          size=16, color=MUTED, align=PP_ALIGN.CENTER)
y = section_close(hd, cy)

OUT = f"{S}/Mars-Fit_포스터.pptx"
prs.save(OUT)
print('저장:', OUT)



# ── HTML 미리보기 ──────────────────────────────────────────────
def emit_html():
    import html as _h
    o = ['<!doctype html><meta charset=utf-8><style>',
         'body{margin:0;background:#999}',
         f'#p{{position:relative;width:{W}px;height:{H}px;overflow:hidden;',
         "font-family:'Noto Sans CJK KR',sans-serif}",
         '#p>*{position:absolute;box-sizing:border-box}',
         '</style><div id=p>']
    for r in OPS:
        if r['k'] == 'box':
            st = [f"left:{r['x']}px", f"top:{r['y']}px",
                  f"width:{r['w']}px", f"height:{r['h']}px",
                  f"background:#{r['fill']}" if r['fill'] else 'background:transparent']
            if r['line']:
                st.append(f"border:{r['lw']*0.353:.2f}px solid #{r['line']}")
            if r['shape'] == 'oval':
                st.append('border-radius:50%')
            elif r['shape'] == 'round':
                st.append(f"border-radius:{min(r['w'],r['h'])*(r['adj'] or .1)*2:.1f}px")
            inner = ''
            if r.get('txt'):
                st += ['display:flex', 'flex-direction:column',
                       'align-items:center', 'justify-content:center',
                       f"font-size:{r['tsize']*0.3528:.2f}px", f"color:#{r['tcolor']}",
                       'font-weight:700' if r['tbold'] else '', 'text-align:center',
                       'padding:0 2px', 'line-height:1.2']
                inner = ''.join(f'<div>{_h.escape(t)}</div>' for t in r['txt'])
            o.append(f'<div style="{";".join(x for x in st if x)}">{inner}</div>')
        elif r['k'] == 'pic':
            o.append(f"<img src=\"file://{r['path']}\" style=\"left:{r['x']}px;"
                     f"top:{r['y']}px;width:{r['w']}px;height:{r['h']}px\">")
        else:
            al = {'LEFT': 'left', 'CENTER': 'center', 'RIGHT': 'right'}[r['align']]
            st = [f"left:{r['x']}px", f"top:{r['y']}px", f"width:{r['w']}px",
                  f"font-size:{r['size']*0.3528:.2f}px", f"color:#{r['color']}",
                  f'text-align:{al}', f"line-height:{r['line_sp']}",
                  'font-weight:700' if r['bold'] else 'font-weight:400',
                  'outline:0.3px dashed rgba(220,0,0,.30)']
            body = []
            for i, parts in enumerate(r['rec']):
                sp = []
                for t, op in parts:
                    ss = []
                    if op.get('bold', r['bold']): ss.append('font-weight:700')
                    if 'color' in op: ss.append(f"color:#{op['color']}")
                    if 'size' in op: ss.append(f"font-size:{op['size']*0.3528:.2f}px")
                    sp.append(f'<span style="{";".join(ss)}">{_h.escape(t)}</span>')
                mt = 0 if i == 0 else r['size'] * r['space'] * 0.3528
                body.append(f'<div style="margin-top:{mt:.1f}px">{"".join(sp)}</div>')
            o.append(f'<div style="{";".join(st)}">{"".join(body)}</div>')
    o.append('</div>')
    io.open(f'{S}/preview.html', 'w', encoding='utf-8').write('\n'.join(o))
    print('미리보기:', f'{S}/preview.html')


emit_html()

print('\n섹션 배치')
cols = {}
for t, x, yy, h in SECTIONS:
    cols.setdefault(x, []).append((t, yy, h))
ok = True
for x, items in sorted(cols.items()):
    print(f'  x={x}mm')
    prev = TOP
    for t, yy, h in items:
        end, flag = yy + h, ''
        if yy < prev - 0.01: flag = '  ← 겹침!'; ok = False
        if end > BOT: flag += f'  ← 넘침({end:.0f}>{BOT})'; ok = False
        print(f'    {t:16} {yy:6.1f} → {end:6.1f}  (h={h:5.1f}){flag}')
        prev = end
    print(f'    남는 높이 {BOT - prev:6.1f}mm')
print('\n판정:', '이상 없음' if ok else '문제 있음')
